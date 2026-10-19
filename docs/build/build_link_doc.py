"""Builds Rig_Sight_CA3_Repository_Link.docx: one page with the repository and demo links.
Usage: python build_link_doc.py <out.docx>"""
import sys

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

REPO_URL, DEMO_URL = "https://github.com/ranjanvrma/rig-sight", "https://rig-sight.streamlit.app/"
ORANGE, MUTED = RGBColor(0xC2, 0x2E, 0x00), RGBColor(0x59, 0x59, 0x59)

doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
sec.left_margin = sec.right_margin = Cm(2.5)
doc.styles["Normal"].font.name = "Calibri"
doc.styles["Normal"].font.size = Pt(11.5)


def para(text="", size=11.5, bold=False, color=None, align=WD_ALIGN_PARAGRAPH.LEFT, after=6):
    p = doc.add_paragraph()
    p.alignment = align
    p.paragraph_format.space_after = Pt(after)
    if text:
        r = p.add_run(text)
        r.font.size, r.bold = Pt(size), bold
        if color:
            r.font.color.rgb = color
    return p


def link(p, url, size=13):
    rid = p.part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
                           is_external=True)
    h = OxmlElement("w:hyperlink"); h.set(qn("r:id"), rid)
    r = OxmlElement("w:r"); rpr = OxmlElement("w:rPr")
    for tag, val in [("w:color", "0563C1"), ("w:u", "single"), ("w:sz", str(int(size * 2)))]:
        e = OxmlElement(tag); e.set(qn("w:val"), val); rpr.append(e)
    b = OxmlElement("w:b"); rpr.append(b)
    r.append(rpr)
    t = OxmlElement("w:t"); t.text = url; r.append(t)
    h.append(r); p._p.append(h)


para("CA-3 · Foundation of Agentic AI", 11, color=ORANGE, after=2)
para("Reading the Rig", 26, bold=True, after=2)
para("Agentic AI for Daily Drilling Reports and non-productive time detection", 12.5, color=MUTED, after=14)
para("Rahul Narkhede (24070126146) and Ranjan Verma (24070126149)", 11.5, after=2)
para("TY B.Tech AI-ML, Division B · Symbiosis Institute of Technology, Pune", 11.5, color=MUTED, after=24)

para("GitHub repository (codebase, PPT and report)", 12, bold=True, after=2)
link(para(after=16), REPO_URL)
para("Live demo (Streamlit Community Cloud)", 12, bold=True, after=2)
link(para(after=24), DEMO_URL)

para("What is in the repository", 12, bold=True, after=4)
for item in [
    "Codebase: app.py (Streamlit UI), graph.py (LangGraph agents and routing), tools.py (Python tools), llm.py "
    "(local model and mock fallback), memory.py (SQLite memory), evaluate.py (evaluation), data/ (synthetic "
    "dataset), tests/ (48 automated tests)",
    "submission/Rig_Sight_CA3_Presentation.pptx: the presentation",
    "submission/Rig_Sight_CA3_Report.docx: the project report",
    "README.md: setup, how to run, results and deployment notes",
]:
    p = doc.add_paragraph(item, style="List Bullet")
    p.paragraph_format.space_after = Pt(3)

para("", after=10)
para("Run locally: pip install -r requirements.txt, then streamlit run app.py. The live demo uses the mock "
     "classifier; the local version can use the qwen2.5:3b model through Ollama.", 10.5, color=MUTED)

doc.core_properties.title = "Reading the Rig - repository link"
doc.core_properties.author = "Ranjan Verma and Rahul Narkhede"
doc.core_properties.last_modified_by = "Ranjan Verma"
doc.save(sys.argv[1])
print("saved", sys.argv[1])
