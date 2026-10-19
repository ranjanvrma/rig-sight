"""Builds Rig_Sight_CA3_Presentation.pptx from the team's Canva deck (PDF export).

Slides whose content is still accurate are kept as high-resolution renders of the original.
Slides with outdated text keep their original artwork, with the old text painted out and new,
editable text placed in the same spot and style. New slides (architecture, live app, results)
use the deck's palette and fonts (Anton titles, Inter body, black / cream / #F63B05).
Usage: python build_deck.py <deck.pdf> <repo> <out.pptx>"""
import io
import json
import sys
from pathlib import Path

import fitz
from PIL import Image, ImageDraw
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Inches, Pt

PDF, REPO, OUT = Path(sys.argv[1]), Path(sys.argv[2]), Path(sys.argv[3])
FIG = REPO / "docs" / "figures"
M = json.loads((REPO / "results" / "metrics_ollama.json").read_text())
MOCK = json.loads((REPO / "results" / "metrics_mock.json").read_text())
G, NC, CB = M["runs"]["Graph with Critic"], M["runs"]["Graph without Critic"], M["runs"]["Single-prompt chatbot"]
N_TESTS = int(sys.argv[4]) if len(sys.argv) > 4 else 48

BLACK, CREAM, ORANGE, WHITE, GREY = (RGBColor(0, 0, 0), RGBColor(0xFF, 0xFA, 0xE5), RGBColor(0xF6, 0x3B, 0x05),
                                     RGBColor(0xFF, 0xFF, 0xFF), RGBColor(0xB9, 0xB2, 0xA0))
TITLE_FONT, BODY_FONT = "Anton", "Inter Light"
W_IN, H_IN = 13.333, 7.5
PT_PER_IN = 72
S = W_IN * PT_PER_IN / 1440          # PDF points -> slide points (0.6667)

doc = fitz.open(PDF)
prs = Presentation()
prs.slide_width, prs.slide_height = Inches(W_IN), Inches(H_IN)
BLANK = prs.slide_layouts[6]


def render(n: int, paint: list[tuple] = ()) -> io.BytesIO:
    """Original slide n as an image; rectangles in PDF points are painted black."""
    zoom = 2400 / 1440
    pix = doc[n - 1].get_pixmap(matrix=fitz.Matrix(zoom, zoom))
    im = Image.frombytes("RGB", (pix.width, pix.height), pix.samples)
    d = ImageDraw.Draw(im)
    for (x0, y0, x1, y1) in paint:
        d.rectangle([x0 * zoom, y0 * zoom, x1 * zoom, y1 * zoom], fill=(0, 0, 0))
    buf = io.BytesIO()
    im.save(buf, "JPEG", quality=88, optimize=True)
    buf.seek(0)
    return buf


def pdf_box(x0, y0, x1, y1):
    return Pt(x0 * S), Pt(y0 * S), Pt((x1 - x0) * S), Pt((y1 - y0) * S)


def text(slide, x, y, w, h, paras, size=12, font=BODY_FONT, color=WHITE, align=PP_ALIGN.LEFT,
         anchor=MSO_ANCHOR.TOP, spacing=None, name=None):
    """paras: list of str or list of (text, {bold, color, font, size}) runs."""
    tb = slide.shapes.add_textbox(x, y, w, h)
    if name:
        tb.name = name
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    for i, p in enumerate(paras):
        para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        para.alignment = align
        if spacing:
            para.space_after = Pt(spacing)
        runs = [(p, {})] if isinstance(p, str) else p
        for t, o in runs:
            r = para.add_run()
            r.text = t
            r.font.name = o.get("font", font)
            r.font.size = Pt(o.get("size", size))
            r.font.bold = o.get("bold", False)
            r.font.color.rgb = o.get("color", color)
    return tb


def bg(slide, color=BLACK):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = color


def image_slide(n, paint=(), notes=""):
    s = prs.slides.add_slide(BLANK)
    bg(s)
    s.shapes.add_picture(render(n, paint), 0, 0, prs.slide_width, prs.slide_height)
    if notes:
        s.notes_slide.notes_text_frame.text = notes
    return s


def kicker(slide, label):
    """The orange Anton label top-right, as on the original content slides."""
    text(slide, Pt(1051 * S), Pt(62 * S), Pt(319 * S), Pt(59 * S), [label], size=27, font=TITLE_FONT,
         color=ORANGE, align=PP_ALIGN.RIGHT)


def logo(slide):
    """Re-use the deck's orange mark: crop it from the original sources slide (black background)."""
    zoom = 4
    r = fitz.Rect(62, 63, 184, 119)
    pix = doc[16].get_pixmap(matrix=fitz.Matrix(zoom, zoom), clip=r)
    buf = io.BytesIO(pix.tobytes("png"))
    slide.shapes.add_picture(buf, Pt(r.x0 * S), Pt(r.y0 * S), Pt(r.width * S), Pt(r.height * S))


# ---------------- slides ----------------

NOTES = {
    1: "Introduce ourselves. Every rig writes a daily report, lost time hides inside it, and our team of AI agents "
       "finds it, checks it and hands it to an engineer to approve.",
    2: "Set the scene with money: one lost hour is about $16,000 of rig time. Operators that track NPT closely cut it "
       "a lot, and big players already pay for agentic AI.",
    3: "Read the three lines. Lines two and three are lost time, but nothing in the report adds them up.",
    4: "NPT has patterns. Today nobody adds it up consistently, so the same failure comes back on the next well.",
    5: "This is the one question the whole project answers.",
    6: "Why not ChatGPT? A chatbot gives one answer with no checking and guesses numbers. We split the job, use "
       "code for the maths and check before a human signs off.",
    7: "Six roles, one shared state. Agents never talk directly; the Orchestrator decides who works next.",
}

for n in range(1, 8):
    image_slide(n, notes=NOTES[n])

# 8 · architecture (the team's original draw.io diagram, as on the Canva deck)
s = prs.slides.add_slide(BLANK)
bg(s, WHITE)
img = Image.open(FIG / "architecture_original.png")
w = prs.slide_width
h = int(w * img.height / img.width)
s.shapes.add_picture(str(FIG / "architecture_original.png"), 0, int((prs.slide_height - h) / 2), w, h)
s.notes_slide.notes_text_frame.text = (
    "Hub and spoke. Orchestrator in the middle, four agents around it, ToolNode for exact maths, incident memory and "
    "reference data at the bottom, engineer review as an interrupt before the final report. Follow arrows 1 to 15. "
    "If asked, the final code differs in small ways: the Planner is rule based, recursion_limit is 50 (worst case "
    "32 steps), the model is qwen2.5:3b, and the Critic also checks gaps, overlaps and unparseable lines.")

# 9 · edge types (accurate as is)
image_slide(9, notes="Normal edges are fixed add_edge calls. Conditional edges come from route() or tools_condition. "
                     "The red retry loop is what makes it agentic.")

# 10 · Python does the maths (first text block updated)
s = image_slide(10, paint=[(40, 688, 700, 760)], notes=(
    "Small local models are fine at labelling a sentence but bad at arithmetic. Every number comes from Python, so "
    "the Critic can check it exactly. We tried letting the model rephrase the summary; it kept the numbers but "
    "attached them to the wrong categories, so the summary is a template by default. No paid API."))
text(s, *pdf_box(43, 692, 690, 760), [
    [("THE MODEL READS: ", {"bold": True}),
     ("a local Ollama model (qwen2.5:3b) only labels each log line with an NPT category and a confidence score. "
      "It does not write the numbers: in testing it swapped which hours belonged to which category. A mock "
      "classifier is the fallback and is always shown as mock.", {})]], size=12)

# 11 · one report (text updated to the real run)
s = image_slide(11, paint=[(1095, 420, 1410, 800)], notes=(
    "This is a real run with qwen2.5:3b. The key moment is the Critic catching 23.5 h and naming the exact gap. "
    "Then the engineer catches what the Critic cannot: a confident mislabel. Both loops are part of the graph."))
text(s, *pdf_box(1102, 410, 1405, 805), [
    [("1 PLAN AND RESEARCH ", {"bold": True}),
     ("Well W-02, Day 14. The Planner makes 9 tasks. Research finds 2,650 m planned against 2,610 m actual and "
      "two past MWD failures on this field.", {})],
    [("2 EXECUTE AND CHECK ", {"bold": True}),
     ("Tools label 5.5 h of tool failure. The Critic sees 23.5 h, names the 20:30-21:00 gap and sends it back. "
      "The retry finds it in REMARKS: 24.0 h, PASS.", {})],
    [("3 ENGINEER DECIDES ", {"bold": True}),
     ("The model also called 'slip and cut' (3 h) NPT. The engineer rejects with 'line 7 is planned', the graph "
      "re-runs: 8.0 h NPT, approved.", {})]], size=11, spacing=6)

# 12, 13 · dataset and demo (accurate as is)
image_slide(12, notes="All synthetic, as the brief asks. We planted every event, so each report has a known answer. "
                      "One trap: 'waited on cement' is planned work, not NPT.")
image_slide(13, notes="Clean day: zero NPT. Tool failure day: right hours and category. W-02 Day 14: the Critic "
                      "catches the 23.5 h log. If the live demo does not fit, use the next slide. Anyone can try it "
                      "at rig-sight.streamlit.app (mock model; the real model runs locally).")

# 14 · the app, live (new)
s = prs.slides.add_slide(BLANK)
bg(s)
logo(s)
kicker(s, "THE APP, LIVE")
text(s, Inches(0.55), Inches(1.15), Inches(7.6), Inches(0.8), ["W-02 · DAY 14, REAL RUN"], size=34,
     font=TITLE_FONT, color=CREAM)
left, right = FIG / "app_w02d14_top.png", FIG / "app_w02d14_approved.png"
gap = Inches(0.35)
lw = Inches(6.4)
im = Image.open(left)
lh = int(lw * im.height / im.width)
rw = prs.slide_width - Inches(0.55) * 2 - lw - gap
im = Image.open(right)
rh = int(rw * im.height / im.width)
top = Inches(2.05)
rx = Inches(0.55) + lw + gap
s.shapes.add_picture(str(left), Inches(0.55), top, lw, lh)
s.shapes.add_picture(str(right), rx, top, rw, rh)
cap_y = top + max(lh, rh) + Inches(0.2)
text(s, Inches(0.55), cap_y, lw, Inches(1.0), [
    [("Run: ", {"bold": True, "color": ORANGE}),
     ("each NPT line labelled with category · hours · model confidence; the 20:30-21:00 line recovered from REMARKS "
      "after the Critic's FAIL; the live agent trace.", {})]], size=12)
text(s, rx, cap_y, rw, Inches(1.0), [
    [("After review: ", {"bold": True, "color": ORANGE}),
     ("the engineer rejected one mislabel (line 7, slip and cut), the graph re-ran: 8.0 h NPT, approved.", {})]],
     size=12)
text(s, Inches(0.55), Inches(6.85), Inches(12.2), Inches(0.4), [
    [("Try it: ", {"bold": True, "color": ORANGE}),
     ("rig-sight.streamlit.app  (hosted on Streamlit Community Cloud, runs on the mock classifier; the real model "
      "runs locally with Ollama)", {})]], size=12)
s.notes_slide.notes_text_frame.text = (
    "Backup for the live demo. Left is the run: the Critic fails the first pass at 23.5 h and the retry recovers the "
    "missing half hour. Right is the final report after the engineer rejected one mislabel and approved.")

# 15 · metrics (method text kept, results text replaced)
s = image_slide(14, paint=[(385, 440, 700, 500)], notes="")
text(s, *pdf_box(390, 438, 700, 520), [
    [("RESULTS (qwen2.5:3b, 15 reports): ", {"bold": True, "color": ORANGE}),
     (f"NPT error {G['npt_hours_mae']:.2f} h per report for our graph vs "
      f"{CB['npt_hours_mae']:.2f} h for a chatbot on the same model. Next slide.", {})]], size=12)
s.notes_slide.notes_text_frame.text = (
    "Every synthetic report has a known answer, so we can score instead of saying 'it looks good'. Same 15 "
    "reports, same model, three configurations.")

# 16 · results (new, native chart + table)
s = prs.slides.add_slide(BLANK)
bg(s)
logo(s)
kicker(s, "RESULTS")
text(s, Inches(0.55), Inches(1.15), Inches(7), Inches(0.8), ["SAME MODEL, 15 REPORTS"], size=34, font=TITLE_FONT,
     color=CREAM)
cd = CategoryChartData()
cd.categories = ["Graph with Critic", "Graph without Critic", "Single-prompt chatbot"]
cd.add_series("NPT hours error per report", (G["npt_hours_mae"], NC["npt_hours_mae"], CB["npt_hours_mae"]))
gf = s.shapes.add_chart(XL_CHART_TYPE.BAR_CLUSTERED, Inches(0.55), Inches(1.95), Inches(5.6), Inches(3.0), cd)
ch = gf.chart
ch.has_legend = False
ch.has_title = True
ch.chart_title.text_frame.text = "Mean NPT hours error per report (lower is better)"
tp = ch.chart_title.text_frame.paragraphs[0]
tp.runs[0].font.size, tp.runs[0].font.color.rgb, tp.runs[0].font.name = Pt(12), CREAM, BODY_FONT
pl = ch.plots[0]
pl.gap_width = 60
pl.has_data_labels = True
dl = pl.data_labels
dl.number_format, dl.number_format_is_linked = '0.00" h"', False
dl.position = XL_LABEL_POSITION.OUTSIDE_END
dl.font.size, dl.font.color.rgb, dl.font.name = Pt(12), CREAM, BODY_FONT
ser = pl.series[0]
for i, col in enumerate([ORANGE, RGBColor(0x9A, 0x4A, 0x30), RGBColor(0x6B, 0x66, 0x5C)]):
    pt = ser.points[i]
    pt.format.fill.solid()
    pt.format.fill.fore_color.rgb = col
cat, val = ch.category_axis, ch.value_axis
cat.reverse_order = True
cat.tick_labels.font.size, cat.tick_labels.font.color.rgb, cat.tick_labels.font.name = Pt(11), CREAM, BODY_FONT
cat.format.line.color.rgb = RGBColor(0x44, 0x44, 0x44)
val.visible = False
val.has_major_gridlines = False
val.maximum_scale = max(5.0, CB["npt_hours_mae"] * 1.25)

rows = [("", "With Critic", "No Critic", "Chatbot"),
        ("Category accuracy", f"{G['category_accuracy']:.0%}", f"{NC['category_accuracy']:.1%}",
         f"{CB['category_accuracy']:.0%}"),
        ("Critic catch rate", "2 / 2", "n/a", "n/a"),
        ("Logs adding to 24 h", G["logs_adding_to_24h"], NC["logs_adding_to_24h"],
         CB["per_repeat"][0]["logs_adding_to_24h"]),
        ("Reports exactly right", G["reports_exact_npt"], NC["reports_exact_npt"],
         CB["per_repeat"][0]["reports_exact_npt"]),
        ("Evidence rate", f"{G['evidence_rate']:.0%}", f"{NC['evidence_rate']:.0%}", f"{CB['evidence_rate']:.0%}")]
tbl = s.shapes.add_table(len(rows), 4, Inches(6.6), Inches(1.95), Inches(6.2), Inches(3.0)).table
tbl.columns[0].width = Inches(2.15)
for c in range(1, 4):
    tbl.columns[c].width = Inches(1.35)
for r, row in enumerate(rows):
    for c, v in enumerate(row):
        cell = tbl.cell(r, c)
        cell.fill.solid()
        cell.fill.fore_color.rgb = RGBColor(0x14, 0x14, 0x14) if r % 2 else RGBColor(0x0A, 0x0A, 0x0A)
        if r == 0:
            cell.fill.fore_color.rgb = RGBColor(0x1E, 0x1E, 0x1E)
        cell.margin_left = cell.margin_right = Inches(0.08)
        tf = cell.text_frame
        tf.text = ""
        run = tf.paragraphs[0].add_run()
        run.text = v
        run.font.size = Pt(12)
        run.font.name = TITLE_FONT if r == 0 else BODY_FONT
        run.font.color.rgb = ORANGE if (c == 1 and r > 0) else CREAM
        tf.paragraphs[0].alignment = PP_ALIGN.LEFT if c == 0 else PP_ALIGN.CENTER
text(s, Inches(0.55), Inches(5.25), Inches(12.2), Inches(1.9), [
    [("What it shows: ", {"bold": True, "color": ORANGE}),
     (f"on the same model, splitting the job and letting Python count cuts the NPT error from {CB['npt_hours_mae']:.2f} h "
      f"to {G['npt_hours_mae']:.2f} h per report. The chatbot cannot even add the log up to 24 h. The Critic caught "
      "both planted errors before any human saw them.", {})],
    [("Read it honestly: ", {"bold": True, "color": ORANGE}),
     (f"'No Critic' has a slightly lower error ({NC['npt_hours_mae']:.2f} h) only because two mistakes cancel on "
      "W-02 Day 14. Our two remaining errors are planned work the model labels NPT with high confidence (slip and "
      "cut, waiting on cement): the engineer fixes each with one reject comment. "
      "Chatbot: latest run; separate runs on other days gave 3.77 h and 4.43 h. "
      f"Mock fallback: {MOCK['runs']['Graph with Critic']['npt_hours_mae']:.2f} h. {N_TESTS} automated tests pass.", {})],
], size=12, spacing=8)
s.notes_slide.notes_text_frame.text = (
    "Same model for all three runs, so the comparison is fair. Point at the chatbot bar first, then ours. Then be "
    "honest about the no-Critic number and the two confident mislabels: that is why the engineer approves every report.")

# 17 · limits + next (limits text updated)
s = image_slide(15, paint=[(62, 556, 520, 640)], notes=(
    "Be honest about limits. It is a PoC on synthetic data. The human stays in the loop because the model makes "
    "confident mistakes and the industry has not settled how much to trust agents."))
text(s, *pdf_box(67, 560, 525, 650), [
    [("LIMITS: ", {"bold": True}),
     ("synthetic data, six NPT categories, one report format. The 3B model labels some planned work as NPT with "
      "high confidence. Not a certified engineering or safety tool.", {})]], size=12)

# 18, 19 · thank you, sources
image_slide(16, notes="Every number computed, every flag sourced, every report approved by a person. Questions.")
image_slide(17, notes="All 11 sources. Source numbers match the small (Source N) tags on the slides.")

prs.core_properties.title = "Reading the Rig - Agentic AI for Daily Drilling Reports (CA-3)"
prs.core_properties.author = "Ranjan Verma and Rahul Narkhede"
prs.core_properties.last_modified_by = "Ranjan Verma"
prs.save(OUT)
print("saved", OUT, len(prs.slides), "slides")
